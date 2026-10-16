# -*- coding: utf-8 -*-
"""
Word 文档渲染器（经验萃取专家 · 三件套生成）

这是一个**渲染器**，不是内容生成器。
内容由你（AI）依据实际素材实时推导后，写成一份 payload JSON，再交给本脚本排版输出。

用法：
    python make_docs.py payload.json [--dir 输出目录] [--md]

参数：
    payload.json   结构见下方 PAYLOAD SCHEMA
    --dir          输出目录，默认 payload 同目录
    --md           强制输出 Markdown（不依赖 python-docx）

降级规则：
    若本机没有 python-docx，自动转为 Markdown 输出，并向用户说明。

PAYLOAD SCHEMA
================================================================
{
  "prefix": "EX-01",                       # 文件名前缀，可选
  "dir": "输出目录（可被 --dir 覆盖）",

  "manual": {                              # ① 操作手册 + 行为复制检查单
    "title": "主标题",
    "subtitle": "副标题",
    "remark": "题注行",
    "info": {"headers": ["项目","内容"], "rows": [["对象","..."]]},
    "checklist_title": "行为复制检查单",
    "checklist_intro": "用法说明",
    "checklist_groups": [
      {"name": "会前（5 条）", "items": ["...", "..."]}
    ],
    "creed": "一句话心法",
    "overview": {"headers": ["阶段","动作","解决什么"], "rows": [["...","...","..."]]},
    "tools": [                             # 每个工具一组 block
      {"h2": "Tool 1　课前调研提纲",
       "blocks": [
         {"type": "para",   "text": "..."},
         {"type": "bullets","items": ["..."]},
         {"type": "table",  "headers": ["a","b"], "rows": [["..."]], "widths": [3.0, 12.6]},
         {"type": "checklist", "items": ["..."]},
         {"type": "keyline","text": "口诀/关键句"},
         {"type": "note",   "text": "灰色提示"},
         {"type": "hr"}
       ]}
    ]
  },

  "case": {                                # ② 案例总结
    "title": "案例总结",
    "subtitle": "...",
    "remark": "...",
    "info": {"headers": [...], "rows": [...]},
    "intro": "给新人看的导语",
    "sections": [
      {"h1": "一、起点", "blocks": [ ...同上... ]}
    ]
  },

  "capability": { ...同上... }             # ③ 可选，勾选才生成
}
================================================================
block.type 支持：h2 / para / bullets / table / checklist / keyline / note / hr
================================================================
"""
import argparse
import json
import os
import re
import sys

CN, EN = '微软雅黑', 'Times New Roman'
BLUE_D = '1F4E79'
BLUE_M = '2E75B6'
GRAY = '595959'
HDR_FILL = 'DDEBF7'

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ---------------------------------------------------------------- Markdown
def render_md(sec):
    """把一个 section（manual / case / capability）渲染成 Markdown 行"""
    lines = []

    def emit_section(sec):
        if sec.get('title'):
            lines.append('# %s' % sec['title'])
            lines.append('')
        if sec.get('subtitle'):
            lines.append('## %s' % sec['subtitle'])
            lines.append('')
        if sec.get('remark'):
            lines.append('> %s' % sec['remark'])
            lines.append('')
        if sec.get('intro'):
            lines.append(sec['intro'])
            lines.append('')
        info = sec.get('info')
        if info:
            lines.append('| %s |' % ' | '.join(info['headers']))
            lines.append('|%s|' % '|'.join(['---'] * len(info['headers'])))
            for r in info['rows']:
                lines.append('| %s |' % ' | '.join(r))
            lines.append('')
        if sec.get('checklist_title'):
            lines.append('## %s' % sec['checklist_title'])
            lines.append('')
            if sec.get('checklist_intro'):
                lines.append(sec['checklist_intro'])
                lines.append('')
            for g in sec.get('checklist_groups', []):
                lines.append('### %s' % g['name'])
                for it in g['items']:
                    lines.append('- [ ] %s' % it)
                lines.append('')
        if sec.get('creed'):
            lines.append('> **%s**' % sec['creed'])
            lines.append('')
        ov = sec.get('overview')
        if ov:
            lines.append('| %s |' % ' | '.join(ov['headers']))
            lines.append('|%s|' % '|'.join(['---'] * len(ov['headers'])))
            for r in ov['rows']:
                lines.append('| %s |' % ' | '.join(r))
            lines.append('')

        def emit_blocks(blocks, base_level=3):
            for b in blocks:
                t = b.get('type')
                if t == 'h2':
                    lines.append('%s %s' % ('#' * base_level, b['text']))
                    lines.append('')
                elif t == 'para':
                    lines.append(b['text'])
                    lines.append('')
                elif t == 'bullets':
                    for i in b['items']:
                        lines.append('- %s' % i)
                    lines.append('')
                elif t == 'table':
                    lines.append('| %s |' % ' | '.join(b['headers']))
                    lines.append('|%s|' % '|'.join(['---'] * len(b['headers'])))
                    for r in b['rows']:
                        lines.append('| %s |' % ' | '.join(r))
                    lines.append('')
                elif t == 'checklist':
                    for i in b['items']:
                        lines.append('- [ ] %s' % i)
                    lines.append('')
                elif t == 'keyline':
                    lines.append('> **%s**' % b['text'])
                    lines.append('')
                elif t == 'note':
                    lines.append('> %s' % b['text'])
                    lines.append('')
                elif t == 'hr':
                    lines.append('---')
                    lines.append('')

        emit_blocks(sec.get('tools_blocks', []))
        for tool in sec.get('tools', []):
            lines.append('### %s' % tool.get('h2', ''))
            lines.append('')
            emit_blocks(tool.get('blocks', []))
        for s in sec.get('sections', []):
            lines.append('## %s' % s.get('h1', ''))
            lines.append('')
            emit_blocks(s.get('blocks', []))

    emit_section(sec)
    lines.append('---')
    lines.append('')

    return lines


def write_md(payload, out_dir, prefix):
    paths = []
    mapping = [('manual', '01_操作手册与行为复制检查单'),
               ('case', '02_案例总结'),
               ('capability', '03_能力与方法价值提炼')]
    for key, name in mapping:
        sec = payload.get(key)
        if not sec:
            continue
        sub = render_md(sec)
        base = name.split('_', 1)[1]
        fname = ('%s_%s.md' % (prefix, name)) if prefix else '%s.md' % name
        p = os.path.join(out_dir, fname)
        with open(p, 'w', encoding='utf-8') as f:
            f.write('\n'.join(sub))
        paths.append(p)
    return paths


# ---------------------------------------------------------------- Word
def _set_run(run, size=12, bold=False, color=None):
    run.font.name = EN
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), CN)
    rfonts.set(qn('w:ascii'), EN)
    rfonts.set(qn('w:hAnsi'), EN)


def _fmt(p, before=0, after=4, indent=0, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.2
    if indent:
        pf.left_indent = Cm(indent)
    if align:
        p.alignment = align


def _new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.4)
    s.top_margin = s.bottom_margin = Cm(2.2)
    st = d.styles['Normal']
    st.font.name = EN
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CN)
    return d


def _title(d, text, size=20, color=BLUE_D, after=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=0):
    p = d.add_paragraph()
    _fmt(p, before=before, after=after, align=align)
    r = p.add_run(text)
    _set_run(r, size=size, bold=True, color=color)


def _table(d, headers, rows, widths=None, size=10.5):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        _fmt(p, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        _set_run(r, size=size, bold=True, color=BLUE_D)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:fill'), HDR_FILL)
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            _fmt(p, after=2)
            r = p.add_run(str(v))
            _set_run(r, size=size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)


def _blocks(d, blocks):
    for b in blocks:
        t = b.get('type')
        if t == 'h2':
            p = d.add_paragraph(); _fmt(p, before=10, after=4)
            r = p.add_run(b['text']); _set_run(r, size=12.5, bold=True, color=BLUE_M)
        elif t == 'para':
            p = d.add_paragraph(); _fmt(p, indent=b.get('indent', 0))
            r = p.add_run(b['text']); _set_run(r, size=b.get('size', 12), bold=b.get('bold', False),
                                               color=GRAY if b.get('gray') else None)
        elif t == 'bullets':
            for i in b['items']:
                p = d.add_paragraph(); _fmt(p, after=3, indent=0.6)
                r = p.add_run('· ' + i); _set_run(r, size=b.get('size', 12))
        elif t == 'table':
            _table(d, b['headers'], b['rows'], b.get('widths'), b.get('size', 10.5))
        elif t == 'checklist':
            for i in b['items']:
                p = d.add_paragraph(); _fmt(p, after=3, indent=0.6)
                r = p.add_run('□  ' + i); _set_run(r, size=11.5)
        elif t == 'keyline':
            p = d.add_paragraph(); _fmt(p, before=6, after=6, indent=0.8)
            r = p.add_run(b['text']); _set_run(r, size=12, bold=True, color=BLUE_M)
        elif t == 'note':
            p = d.add_paragraph(); _fmt(p, after=4, indent=0.5)
            r = p.add_run(b['text']); _set_run(r, size=11.5, color=GRAY)
        elif t == 'hr':
            p = d.add_paragraph(); _fmt(p, after=8)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bot = OxmlElement('w:bottom')
            bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6'); bot.set(qn('w:color'), 'BDD7EE')
            pbdr.append(bot); pPr.append(pbdr)
        elif t == 'pagebreak':
            d.add_page_break()


def _build_section(d, sec):
    if sec.get('title'):
        _title(d, sec['title'], size=22, before=20, after=2)
    if sec.get('subtitle'):
        _title(d, sec['subtitle'], size=15, color=BLUE_M, after=6)
    if sec.get('remark'):
        _title(d, sec['remark'], size=11.5, color=GRAY, after=20)
    if sec.get('info'):
        _table(d, sec['info']['headers'], sec['info']['rows'], [3.0, 12.6], 11)
    if sec.get('intro'):
        p = d.add_paragraph(); _fmt(p, after=6)
        r = p.add_run(sec['intro']); _set_run(r, size=11.5, color=GRAY)
    if sec.get('checklist_title'):
        d.add_page_break()
        p = d.add_paragraph(); _fmt(p, before=10, after=6)
        r = p.add_run(sec['checklist_title']); _set_run(r, size=14, bold=True, color=BLUE_D)
        if sec.get('checklist_intro'):
            p = d.add_paragraph(); _fmt(p, after=8)
            r = p.add_run(sec['checklist_intro']); _set_run(r, size=11.5, color=GRAY)
        for g in sec.get('checklist_groups', []):
            p = d.add_paragraph(); _fmt(p, before=8, after=4)
            r = p.add_run(g['name']); _set_run(r, size=12.5, bold=True, color=BLUE_M)
            for it in g['items']:
                p = d.add_paragraph(); _fmt(p, after=3, indent=0.6)
                r = p.add_run('□  ' + it); _set_run(r, size=11.5)
    if sec.get('creed'):
        p = d.add_paragraph(); _fmt(p, before=8, after=6, indent=0.8)
        r = p.add_run(sec['creed']); _set_run(r, size=12, bold=True, color=BLUE_M)
    if sec.get('overview'):
        d.add_page_break()
        _blocks(d, [{'type': 'para', 'text': '方法总览', 'bold': True, 'size': 13}])
        _blocks(d, [{'type': 'table', 'headers': sec['overview']['headers'],
                     'rows': sec['overview']['rows'], 'widths': [1.8, 7.6, 6.2]}])
    _blocks(d, sec.get('tools_blocks', []))
    if sec.get('tools'):
        d.add_page_break()
    for tool in sec.get('tools', []):
        if tool.get('h2'):
            p = d.add_paragraph(); _fmt(p, before=10, after=4)
            r = p.add_run(tool['h2']); _set_run(r, size=12.5, bold=True, color=BLUE_M)
        _blocks(d, tool.get('blocks', []))
    for s in sec.get('sections', []):
        if s.get('h1'):
            p = d.add_paragraph(); _fmt(p, before=14, after=6)
            r = p.add_run(s['h1']); _set_run(r, size=14, bold=True, color=BLUE_D)
        _blocks(d, s.get('blocks', []))


def write_docx(payload, out_dir, prefix):
    paths = []
    mapping = [('manual', '01_操作手册与行为复制检查单'),
               ('case', '02_案例总结'),
               ('capability', '03_能力与方法价值提炼')]
    for key, name in mapping:
        sec = payload.get(key)
        if not sec:
            continue
        d = _new_doc()
        _build_section(d, sec)
        fname = ('%s_%s.docx' % (prefix, name)) if prefix else '%s.docx' % name
        p = os.path.join(out_dir, fname)
        d.save(p)
        paths.append(p)
    return paths


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description='经验萃取三件套渲染器')
    ap.add_argument('payload', help='结构描述 JSON')
    ap.add_argument('--dir', help='输出目录（覆盖 payload 中的 dir）')
    ap.add_argument('--md', action='store_true', help='强制 Markdown 输出')
    args = ap.parse_args()

    with open(args.payload, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    out_dir = args.dir or payload.get('dir') or os.path.dirname(os.path.abspath(args.payload)) or os.getcwd()
    os.makedirs(out_dir, exist_ok=True)
    prefix = payload.get('prefix', '')

    if HAS_DOCX and not args.md:
        paths = write_docx(payload, out_dir, prefix)
        mode = 'Word (.docx)'
    else:
        if args.md:
            mode = 'Markdown (.md，用户指定)'
        else:
            print('[降级] 未检测到 python-docx，改为输出 Markdown。安装方法：pip install python-docx',
                  file=sys.stderr)
            mode = 'Markdown (.md，因缺少 python-docx 降级)'
        paths = write_md(payload, out_dir, prefix)

    print('输出格式：%s' % mode)
    for p in paths:
        print('  -> %s' % p)


if __name__ == '__main__':
    main()
